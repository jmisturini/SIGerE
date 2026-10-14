"""Testes do módulo Adicional Noturno - Professores (RH).

Adaptação web do gerador local "Gerador-Ad-Noturno-SENAC-FL": formulário
com as colunas da planilha "Adicional Noturno Professores" e exportação
da planilha de pagamento no modelo planilha_base_an.xlsx. Regras cobertas:
validação dos campos, critérios do gerador (adicional "Sim" com minutos
maiores que 0), conversão de minutos decimais, ordenação por unidade (UO),
escopo por unidade, permissões (an:*) e o toggle Financeiro da unidade.
"""
import os
import tempfile
import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import create_app
from app.config import Config
from app.extensions import db
from app.models import (Permission, Role, TeacherNightShiftPay, Unity, User)

EMAIL = 'gestor@escola.edu'
PASSWORD = 'SenhaForte123'


class TestConfig(Config):
    SECRET_KEY = 'chave-de-teste-nao-usar-o-valor-dev'
    TESTING = True
    WTF_CSRF_ENABLED = False
    RATELIMIT_ENABLED = False


class AdicionalNoturnoTestCase(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix='.db')
        os.close(fd)
        TestConfig.SQLALCHEMY_DATABASE_URI = 'sqlite:///' + self.db_path.replace('\\', '/')

        self.app = create_app(TestConfig)
        self.app.config['SESSION_COOKIE_SECURE'] = False
        self.client = self.app.test_client()

        with self.app.app_context():
            db.create_all()
            self.unity = Unity(name='Unidade Teste', code='UT')
            db.session.add(self.unity)
            db.session.flush()

            perms = [Permission(code='an:read', module='an', action='read'),
                     Permission(code='an:create', module='an', action='create'),
                     Permission(code='an:edit', module='an', action='edit'),
                     Permission(code='an:delete', module='an', action='delete'),
                     Permission(code='an:export', module='an', action='export')]
            db.session.add_all(perms)
            role = Role(name='gestor-teste', label='Gestor Teste', permissions=perms)
            db.session.add(role)
            db.session.flush()

            gestor = User(
                email=EMAIL, full_name='Gestor Teste',
                role='room', profile_type='employee', unities=[self.unity],
                role_id=role.id, force_password_change=False, is_active_user=True,
            )
            gestor.set_password(PASSWORD)
            db.session.add(gestor)
            db.session.flush()

            # Professor do cadastro com matrícula — alimenta o select do
            # formulário e o auto-preenchimento.
            self.professor = User(
                email='professor@escola.edu', full_name='Ana Paula Souza',
                role='room', profile_type='teacher', unities=[self.unity],
                registration='13164', force_password_change=False,
                is_active_user=True,
            )
            self.professor.set_password(PASSWORD)
            db.session.add(self.professor)
            db.session.commit()
            self.gestor_id = gestor.id
            self.professor_id = self.professor.id
            self.unity_id = self.unity.id

        response = self.client.post('/login', data={'email': EMAIL, 'password': PASSWORD},
                                    follow_redirects=True)
        self.assertEqual(response.status_code, 200)

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        for suffix in ('', '-wal', '-shm'):
            path = self.db_path + suffix
            if os.path.exists(path):
                os.remove(path)

    def _adicionar(self, teacher='manual', matricula='13164',
                   nome='Adriano de Melo', adicional='Sim',
                   horas='58,91', minutos='160', uo='2', mes=None):
        data = {'teacher': teacher, 'registration': matricula,
                'full_name': nome, 'entitled': adicional,
                'class_hours': horas, 'night_minutes': minutos,
                'uo_code': uo}
        if mes is not None:
            data['month_base'] = mes
        return self.client.post('/adicional-noturno/adicionar',
                                data=data, follow_redirects=True)

    def _id_unico(self):
        with self.app.app_context():
            return TeacherNightShiftPay.query.one().id

    def _criar_registro(self, **kwargs):
        """Registro direto no banco para os testes de exportação."""
        dados = dict(unity_id=self.unity_id, month_base='2026-10',
                     registration='13164', full_name='Adriano de Melo',
                     entitled=True, class_hours=58.91, night_minutes=160,
                     uo_code=2, created_by_id=self.gestor_id)
        dados.update(kwargs)
        registro = TeacherNightShiftPay(**dados)
        db.session.add(registro)
        return registro

    # ---------- Minutos decimais (conversão do gerador) ----------

    def test_minutos_decimais_conversao(self):
        """A conversão da planilha de origem: h×100+m volta a minutos
        decimais ((h×100+m)×60/100) — 2h40 → 144 e 0h40 → 24."""
        self.assertEqual(TeacherNightShiftPay(night_minutes=160).decimal_minutes, 144)
        self.assertEqual(TeacherNightShiftPay(night_minutes=40).decimal_minutes, 24)
        self.assertEqual(TeacherNightShiftPay(night_minutes=90).decimal_minutes, 78)
        self.assertEqual(TeacherNightShiftPay(night_minutes=0).decimal_minutes, 0)
        self.assertEqual(TeacherNightShiftPay().decimal_minutes, 0)

    def test_exportavel_exige_sim_e_minutos(self):
        """Critérios do gerador original: adicional "Sim" com minutos > 0."""
        self.assertTrue(TeacherNightShiftPay(entitled=True, night_minutes=160).is_exportable)
        self.assertFalse(TeacherNightShiftPay(entitled=True, night_minutes=0).is_exportable)
        self.assertFalse(TeacherNightShiftPay(entitled=False, night_minutes=160).is_exportable)
        self.assertFalse(TeacherNightShiftPay(entitled=False, night_minutes=None).is_exportable)

    # ---------- Listagem e formulário ----------

    def test_pagina_mostra_form_listagem_e_menu(self):
        response = self.client.get('/adicional-noturno/')
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('Adicional Noturno', html)
        self.assertIn('Matrícula', html)
        self.assertIn('Horas/Aula', html)
        self.assertIn('Nenhum lançamento adicionado', html)
        # item do menu RH no sidebar
        self.assertIn('/adicional-noturno', html)
        # ações do cabeçalho: exportação e limpeza geral
        self.assertIn('/adicional-noturno/exportar', html)
        self.assertIn('/adicional-noturno/excluir-todos', html)
        # referência dos códigos UO
        self.assertIn('Florianópolis', html)

    def test_adicionar_lancamento(self):
        response = self._adicionar()
        self.assertIn('Lançamento adicionado', response.get_data(as_text=True))
        with self.app.app_context():
            registro = TeacherNightShiftPay.query.one()
            self.assertEqual(registro.registration, '13164')
            self.assertEqual(registro.full_name, 'Adriano de Melo')
            self.assertTrue(registro.entitled)
            self.assertEqual(float(registro.class_hours), 58.91)
            self.assertEqual(registro.night_minutes, 160)
            self.assertEqual(registro.uo_code, 2)
            self.assertEqual(registro.unity_id, self.unity_id)
            self.assertIsNotNone(registro.created_by_id)
            self.assertIsNotNone(registro.month_base)
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertIn('Adriano de Melo', html)
        self.assertIn('58.91', html)  # horas/aula na listagem
        self.assertIn('view-an-', html)
        self.assertIn('Excluir este lançamento', html)

    def test_editar_lancamento(self):
        # Matrícula que não casa com nenhum professor do cadastro: a edição
        # abre no modo manual, com o nome gravado no campo de digitação.
        self._adicionar(matricula='88888')
        registro_id = self._id_unico()

        response = self.client.get(f'/adicional-noturno/{registro_id}/editar')
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('Editar Lançamento', html)
        self.assertIn('value="88888"', html)
        self.assertIn('value="58,91"', html)
        self.assertIn('<option selected value="manual">', html)
        self.assertIn('value="Adriano de Melo"', html)

        response = self.client.post(f'/adicional-noturno/{registro_id}/editar',
                                    data={'teacher': 'manual',
                                          'registration': '99999',
                                          'full_name': 'Carlize Bortoli',
                                          'entitled': 'Não',
                                          'class_hours': '41,13',
                                          'night_minutes': '40',
                                          'uo_code': '4',
                                          'month_base': '2026-11'},
                                    follow_redirects=True)
        self.assertIn('Alteração realizada', response.get_data(as_text=True))
        with self.app.app_context():
            registro = db.session.get(TeacherNightShiftPay, registro_id)
            self.assertEqual((registro.registration, registro.full_name,
                              registro.entitled, registro.uo_code,
                              registro.month_base),
                             ('99999', 'Carlize Bortoli', False, 4, '2026-11'))

    def test_editar_preseleciona_professor_do_cadastro(self):
        """Lançamento com a matrícula de um professor do cadastro abre a
        edição com ele pré-selecionado no select."""
        self._adicionar()  # manual, com a matrícula 13164 do professor
        registro_id = self._id_unico()
        html = self.client.get(
            f'/adicional-noturno/{registro_id}/editar').get_data(as_text=True)
        self.assertIn(f'<option selected value="{self.professor_id}">', html)

        # Salvando sem mexer no select: os dados continuam vindo da conta.
        response = self.client.post(f'/adicional-noturno/{registro_id}/editar',
                                    data={'teacher': str(self.professor_id),
                                          'registration': '13164',
                                          'entitled': 'Sim',
                                          'class_hours': '58,91',
                                          'night_minutes': '200',
                                          'uo_code': '2',
                                          'month_base': '2026-10'},
                                    follow_redirects=True)
        self.assertIn('Alteração realizada', response.get_data(as_text=True))
        with self.app.app_context():
            registro = db.session.get(TeacherNightShiftPay, registro_id)
            self.assertEqual((registro.full_name, registro.registration),
                             ('Ana Paula Souza', '13164'))

    def test_excluir_lancamento(self):
        self._adicionar()
        registro_id = self._id_unico()
        response = self.client.post(f'/adicional-noturno/{registro_id}/excluir',
                                    follow_redirects=True)
        self.assertIn('excluído', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 0)

    def test_validacoes_recusadas(self):
        casos = [
            ({'matricula': ''}, 'Informe a matrícula'),
            ({'nome': ''}, 'Informe o nome do professor'),
            ({'minutos': ''}, 'Informe os minutos noturnos'),
            ({'minutos': '2000'}, 'entre 0 e 1440'),
            ({'uo': ''}, 'Informe o código da unidade'),
            ({'horas': '58,9,1'}, 'Formato inválido'),
        ]
        for extra, mensagem in casos:
            dados = dict(teacher='manual', matricula='13164', nome='Adriano de Melo',
                         adicional='Sim', horas='58,91', minutos='160', uo='2')
            dados.update(extra)
            response = self._adicionar(**dados)
            self.assertIn(mensagem, response.get_data(as_text=True),
                          msg=f'caso {extra}')
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 0)

    def test_selecionar_professor_do_cadastro(self):
        """Professor selecionado no cadastro: nome e matrícula vêm da CONTA,
        mesmo que a página envie nome/matrícula em branco."""
        response = self._adicionar(teacher=str(self.professor_id),
                                   matricula='', nome='')
        self.assertIn('Lançamento adicionado', response.get_data(as_text=True))
        with self.app.app_context():
            registro = TeacherNightShiftPay.query.one()
            self.assertEqual(registro.full_name, 'Ana Paula Souza')
            self.assertEqual(registro.registration, '13164')
        # O select da listagem oferece o professor do cadastro.
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertIn('Ana Paula Souza', html)
        self.assertIn('Outro (digitar nome e matrícula)', html)

    def test_professor_sem_matricula_no_cadastro(self):
        with self.app.app_context():
            professor = User(
                email='semmatricula@escola.edu', full_name='Sem Matrícula',
                role='room', profile_type='teacher', unities=[self.unity],
                force_password_change=False, is_active_user=True,
            )
            professor.set_password(PASSWORD)
            db.session.add(professor)
            db.session.commit()
            professor_id = professor.id
        response = self._adicionar(teacher=str(professor_id), matricula='')
        self.assertIn('não tem matrícula no cadastro',
                      response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 0)

    def test_professor_nao_selecionado(self):
        response = self._adicionar(teacher='')
        self.assertIn('Selecione o professor', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 0)

    def test_minutos_zero_aceito_sem_exportar(self):
        """Sim com minutos 0 é lançamento válido, mas fica fora da planilha —
        sem elegíveis, a exportação redireciona de volta."""
        self._adicionar(minutos='0')
        with self.app.app_context():
            registro = TeacherNightShiftPay.query.one()
            self.assertFalse(registro.is_exportable)
        response = self._exportar()
        self.assertEqual(response.status_code, 302)

    def test_escopo_unidade(self):
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            self._criar_registro(unity_id=outra.id)
            db.session.commit()
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertIn('Nenhum lançamento adicionado', html)

    def test_sem_permissao_vira_403(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = []
            db.session.commit()
        self.assertEqual(self.client.get('/adicional-noturno/').status_code, 403)
        self.assertEqual(self._adicionar().status_code, 403)

    def test_leitura_sem_criacao_esconde_form(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions if p.code == 'an:read']
            db.session.commit()
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertNotIn('an.adicionar', html)
        self.assertIn('Nenhum lançamento adicionado', html)
        self.assertEqual(self._adicionar().status_code, 403)

    def test_editar_excluir_sem_permissao_vira_403(self):
        self._adicionar()
        registro_id = self._id_unico()
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions
                                if p.code in ('an:read', 'an:create')]
            db.session.commit()
        self.assertEqual(
            self.client.get(f'/adicional-noturno/{registro_id}/editar').status_code, 403)
        self.assertEqual(
            self.client.post(f'/adicional-noturno/{registro_id}/editar',
                             data={'registration': '13164',
                                   'full_name': 'Adriano de Melo',
                                   'entitled': 'Sim', 'night_minutes': '160',
                                   'uo_code': '2'}).status_code, 403)
        self.assertEqual(
            self.client.post(f'/adicional-noturno/{registro_id}/excluir').status_code,
            403)
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 1)

    def test_modulo_financeiro_desligado_vira_403(self):
        with self.app.app_context():
            unity = db.session.get(Unity, self.unity_id)
            unity.finance_enabled = False
            db.session.commit()
        self.assertEqual(self.client.get('/adicional-noturno/').status_code, 403)
        self.assertEqual(self._adicionar().status_code, 403)

    # ---------- Filtro por mês ----------

    def test_filtro_por_mes(self):
        self._adicionar(nome='Professor Agosto', mes='2026-08')
        self._adicionar(nome='Professor Setembro', mes='2026-09')

        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertIn('Professor Agosto', html)
        self.assertIn('Professor Setembro', html)

        filtrado = self.client.get(
            '/adicional-noturno/?month_base=2026-08').get_data(as_text=True)
        self.assertIn('Professor Agosto', filtrado)
        self.assertNotIn('Professor Setembro', filtrado)
        self.assertIn('<option value="2026-08" selected>', filtrado)

    # ---------- Limpar todos os lançamentos ----------

    def test_limpar_todos_lancamentos(self):
        self._adicionar(nome='Professor A')
        self._adicionar(nome='Professor B')

        response = self.client.post('/adicional-noturno/excluir-todos',
                                    follow_redirects=True)
        self.assertIn('2 lançamento(s) do Adicional Noturno removido(s)',
                      response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 0)

    def test_limpar_todos_escopo_unidade(self):
        self._adicionar(nome='Professor A')
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            self._criar_registro(unity_id=outra.id)
            db.session.commit()
            outra_id = outra.id

        self.client.post('/adicional-noturno/excluir-todos')
        with self.app.app_context():
            restante = TeacherNightShiftPay.query.one()
            self.assertEqual(restante.unity_id, outra_id)

    def test_limpar_todos_sem_permissao_vira_403(self):
        self._adicionar(nome='Professor A')
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions if p.code == 'an:read']
            db.session.commit()

        # an:read sozinho vê a página, mas não tem o botão nem a rota.
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertNotIn('/adicional-noturno/excluir-todos', html)
        self.assertEqual(
            self.client.post('/adicional-noturno/excluir-todos').status_code, 403)
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 1)

    # ---------- Exportação Excel ----------

    def _exportar(self, query=''):
        return self.client.get('/adicional-noturno/exportar' + query)

    def _linhas_planilha(self, response):
        wb = load_workbook(BytesIO(response.data))
        ws = wb['Adicional Noturno']
        return [[ws.cell(row=r, column=c).value for c in (1, 2, 3, 7, 11)]
                for r in range(5, 12)]

    def test_exportar_planilha_pagamento(self):
        self._adicionar(nome='Adriano de Melo', minutos='160', uo='2')

        response = self._exportar()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('Tabela Adicional Noturno Professores',
                      response.headers['Content-Disposition'])

        linhas = self._linhas_planilha(response)
        # Matrícula numérica, nome, h/a, minutos decimais e UO na linha 5.
        self.assertEqual(linhas[0], [13164, 'Adriano de Melo', 58.91, 144, 2])
        # Linhas seguintes ficam com as fórmulas do modelo, sem dados
        # (a coluna K traz o padrão 2 do próprio modelo).
        self.assertEqual(linhas[1], [None, None, None, None, 2])

    def test_exportar_critérios_do_gerador(self):
        """Só entra na planilha quem tem Sim + minutos > 0."""
        self._adicionar(matricula='13164', nome='Elegível', minutos='160', uo='2')
        self._adicionar(matricula='22222', nome='Zero Minutos',
                        minutos='0', uo='2')
        self._adicionar(matricula='33333', nome='Sem Adicional',
                        adicional='Não', minutos='90', uo='2')

        linhas = self._linhas_planilha(self._exportar())
        self.assertEqual(linhas[0][1], 'Elegível')
        nomes = {linha[1] for linha in linhas[1:]}
        self.assertNotIn('Zero Minutos', nomes)
        self.assertNotIn('Sem Adicional', nomes)
        self.assertIsNone(linhas[1][0])

    def test_exportar_ordenado_por_unidade(self):
        """Ordenação do gerador original: pelo código da unidade, mantendo
        dentro de cada unidade a ordem de lançamento."""
        self._adicionar(matricula='2', nome='Z lançado primeiro, Unidade 4', uo='4')
        self._adicionar(matricula='1', nome='Ana Unidade 2', uo='2')
        self._adicionar(matricula='3', nome='Bruno Unidade 2', uo='2')

        linhas = self._linhas_planilha(self._exportar())
        self.assertEqual([linha[1] for linha in linhas[:3]],
                         ['Ana Unidade 2', 'Bruno Unidade 2',
                          'Z lançado primeiro, Unidade 4'])
        self.assertEqual([linha[4] for linha in linhas[:3]], [2, 2, 4])

    def test_exportar_respeita_filtro_por_mes(self):
        self._adicionar(nome='Professor Agosto', mes='2026-08')
        self._adicionar(matricula='55555', nome='Professor Setembro',
                        mes='2026-09')

        linhas = self._linhas_planilha(self._exportar('?month_base=2026-09'))
        self.assertEqual(linhas[0][1], 'Professor Setembro')
        self.assertIsNone(linhas[1][0])

    def test_exportar_escopo_unidade(self):
        self._adicionar(nome='Professor Local')
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            self._criar_registro(unity_id=outra.id, full_name='Professor Fora')
            db.session.commit()

        linhas = self._linhas_planilha(self._exportar())
        self.assertEqual(linhas[0][1], 'Professor Local')
        self.assertIsNone(linhas[1][0])

    def test_exportar_sem_elegiveis_redireciona(self):
        response = self._exportar()
        self.assertEqual(response.status_code, 302)
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertIn('Nenhum lançamento elegível para exportar', html)

    def test_exportar_sem_permissao_vira_403(self):
        self._adicionar(nome='Professor A')
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions
                                if p.code != 'an:export']
            db.session.commit()
        self.assertEqual(self._exportar().status_code, 403)

    def test_exportar_preserva_modelo(self):
        """A exportação preenche apenas os dados — as fórmulas do modelo
        fornecido, a referência de minutos (G2), o cabeçalho, a aba oculta e
        o recálculo na abertura permanecem intactos."""
        self._adicionar(nome='Adriano de Melo')

        wb = load_workbook(BytesIO(self._exportar().data))
        ws = wb['Adicional Noturno']
        # Fórmulas de pagamento das linhas de dados e a base de conversão.
        self.assertEqual(ws['E5'].value, '=PRODUCT(C5,D5)')
        self.assertEqual(ws['H5'].value, '=G5/$G$2')
        self.assertEqual(ws['I5'].value, '=PRODUCT(H5*F5)')
        self.assertEqual(ws['J5'].value, '=PRODUCT(E5*I5)')
        self.assertEqual(ws['G2'].value, 60)
        self.assertEqual(ws['A1'].value,
                         'TABELA PARA PAGAMENTO ADICIONAL NOTURNO PROFESSOR')
        # Aba oculta continua referenciando as linhas da planilha.
        oculta = wb['Não mexer nesta planilha']
        self.assertEqual(oculta['B1'].value, "='Adicional Noturno'!K5")
        self.assertEqual(oculta['G1'].value, "='Adicional Noturno'!J5")
        # Fórmulas além dos dados lançados permanecem (linha 100).
        self.assertEqual(ws['E100'].value, '=PRODUCT(C100,D100)')
        # Excel recalcula as fórmulas ao abrir o arquivo gerado.
        self.assertTrue(wb.calculation.fullCalcOnLoad)

    def test_exportar_preenche_resultados_em_cache(self):
        """As colunas calculadas mostram valores ao abrir o arquivo (mesmo
        sem recálculo, como no Modo de Exibição Protegido): o openpyxl
        descarta o cache que o Excel guarda junto das fórmulas, e a
        exportação o repõe — sem alterar as fórmulas."""
        self._adicionar(nome='Adriano de Melo')  # 58,91 h/a e 160 min (144 dec.)

        valores = load_workbook(BytesIO(self._exportar().data),
                                data_only=True)
        ws = valores['Adicional Noturno']
        # E = C*D (58,91 × 0,2), H = G/G2 (144/60), I = H*F, J = E*I.
        self.assertAlmostEqual(ws['E5'].value, 58.91 * 0.2)
        self.assertAlmostEqual(ws['H5'].value, 2.4)
        self.assertAlmostEqual(ws['I5'].value, 2.4)
        self.assertAlmostEqual(ws['J5'].value, 58.91 * 0.2 * 2.4)
        # Linha sem lançamento fica com 0, como na planilha original do Excel.
        self.assertAlmostEqual(ws['E6'].value, 0)
        self.assertAlmostEqual(ws['J6'].value, 0)
        # Aba oculta: referências em cache para a linha de dados.
        oculta = valores['Não mexer nesta planilha']
        self.assertAlmostEqual(oculta['B1'].value, 2)          # UO
        self.assertAlmostEqual(oculta['C1'].value, 13164)      # matrícula
        self.assertAlmostEqual(oculta['E1'].value, 2.4)        # total h/m
        self.assertAlmostEqual(oculta['G1'].value, 58.91 * 0.2 * 2.4)

    def test_exportar_cache_aceita_matricula_texto(self):
        """Matrícula alfanumérica (professor do cadastro) também é reposta no
        cache da aba oculta, como texto."""
        self._adicionar(matricula='MAT-2024-0001', nome='Ana Paula Souza')

        oculta = load_workbook(BytesIO(self._exportar().data),
                               data_only=True)['Não mexer nesta planilha']
        self.assertEqual(oculta['C1'].value, 'MAT-2024-0001')

    def test_exportar_alarga_colunas_com_dados(self):
        """As colunas com conteúdo maior que a largura do modelo são alargadas
        (a coluna E do 20% exibiria '###'); as outras mantêm a largura
        original do modelo."""
        self._adicionar(nome='Adriano de Melo')

        ws = load_workbook(BytesIO(self._exportar().data))['Adicional Noturno']
        # E no modelo tem 4,71 (só cabe '0,00'); com 11,78 precisa de mais.
        self.assertGreaterEqual(ws.column_dimensions['E'].width, 7.5)
        # A também alarga para matrícula alfanumérica; lançamento novo para isso.
        self._limpar_e_lancar_matricula_texto()
        ws = load_workbook(BytesIO(self._exportar().data))['Adicional Noturno']
        self.assertGreaterEqual(ws.column_dimensions['A'].width, 15)
        # B é larga no modelo (29,43) e não precisa mudar para nomes curtos;
        # D nunca recebe dados e permanece como no modelo.
        self.assertAlmostEqual(ws.column_dimensions['B'].width, 29.43, places=2)
        self.assertAlmostEqual(ws.column_dimensions['D'].width, 1.86, places=2)

    def _limpar_e_lancar_matricula_texto(self):
        self.client.post('/adicional-noturno/excluir-todos')
        self.client.post('/adicional-noturno/adicionar',
                         data={'teacher': 'manual', 'registration': 'MAT-2024-0001',
                               'full_name': 'Ana Paula Souza', 'entitled': 'Sim',
                               'class_hours': '58,91', 'night_minutes': '160',
                               'uo_code': '2'}, follow_redirects=True)

    # ---------- Importar planilha ----------

    CABECALHO_MODELO = ['Matrícula', 'Professor', 'Adicional Noturno',
                        'Horas/Aula', 'Adicional Noturno (Minutos)',
                        'Adicional Noturno (Horas)', 'Conversão de Horas',
                        'Minutos Decimais', 'Unidade']

    def _planilha_bytes(self, linhas, cabecalho=True):
        """Bytes de um .xlsx no modelo 'Adicional Noturno Professores'."""
        wb = Workbook()
        ws = wb.active
        if cabecalho:
            ws.append(self.CABECALHO_MODELO)
        for linha in linhas:
            ws.append(linha)
        saida = BytesIO()
        wb.save(saida)
        saida.seek(0)
        return saida

    def _importar(self, planilha, nome='adicional.xlsx', mes=None):
        data = {'planilha': (planilha, nome)}
        if mes:
            data['month_base'] = mes
        return self.client.post('/adicional-noturno/importar', data=data,
                                content_type='multipart/form-data',
                                follow_redirects=True)

    def test_caixa_importacao_na_pagina(self):
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertIn('Importar Planilha', html)
        self.assertIn('/adicional-noturno/importar', html)
        self.assertIn('Mês dos Lançamentos', html)

    def test_caixa_importacao_sem_permissao_de_criacao(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions
                                if p.code == 'an:read']
            db.session.commit()
        html = self.client.get('/adicional-noturno/').get_data(as_text=True)
        self.assertNotIn('/adicional-noturno/importar', html)

    def test_importar_planilha_do_gerador(self):
        planilha = self._planilha_bytes([
            [13164, 'Adriano de Melo', 'Sim', 58.91, 160, None, None, None, 2],
            [14501, 'Alan Giovani da Silveira', 'Não', 58.91, 0,
             None, None, None, 2],
            [None, 'Sem matrícula', 'Sim', 30, 40, None, None, None, 2],
        ])
        response = self._importar(planilha, mes='2026-10')
        html = response.get_data(as_text=True)
        self.assertIn('2 lançamento(s) importado(s) para Outubro/2026', html)
        self.assertIn('1 linha(s) sem matrícula ou professor foram ignoradas',
                      html)
        with self.app.app_context():
            registros = {r.registration: r
                         for r in TeacherNightShiftPay.query.all()}
            self.assertEqual(set(registros), {'13164', '14501'})
            self.assertTrue(registros['13164'].entitled)
            self.assertEqual(registros['13164'].night_minutes, 160)
            self.assertEqual(float(registros['13164'].class_hours), 58.91)
            self.assertEqual(registros['13164'].uo_code, 2)
            self.assertFalse(registros['14501'].entitled)
            self.assertEqual(registros['14501'].night_minutes, 0)
        # A importação abre no mês escolhido, com o filtro aplicado.
        self.assertIn('<option value="2026-10" selected>', html)

    def test_importar_nao_duplica_e_lista_mantidos(self):
        self._adicionar()  # matrícula 13164 já lançada no mês corrente
        planilha = self._planilha_bytes([
            [13164, 'Adriano de Melo', 'Sim', 58.91, 160, None, None, None, 2],
            [7616, 'Ronaldo Marcial Pasquini', 'Sim', 58.91, 120,
             None, None, None, 2],
        ])
        response = self._importar(planilha)
        html = response.get_data(as_text=True)
        self.assertIn('1 lançamento(s) importado(s)', html)
        self.assertIn('1 professor(es) já possuíam lançamento no mês e '
                      'foram mantidos: Adriano de Melo', html)
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 2)

    def test_importar_meses_diferentes(self):
        for mes in ('2026-08', '2026-09'):
            planilha = self._planilha_bytes([
                [13164, 'Adriano de Melo', 'Sim', 58.91, 160,
                 None, None, None, 2],
            ])
            self._importar(planilha, mes=mes)
        with self.app.app_context():
            self.assertEqual(sorted(r.month_base for r in
                                    TeacherNightShiftPay.query.all()),
                             ['2026-08', '2026-09'])

    def test_importar_sem_arquivo(self):
        response = self.client.post('/adicional-noturno/importar', data={},
                                    follow_redirects=True)
        self.assertIn('Selecione a planilha (.xlsx) para importar',
                      response.get_data(as_text=True))

    def test_importar_formato_invalido(self):
        response = self._importar(BytesIO(b'nao sou xlsx'), nome='planilha.txt')
        self.assertIn('Formato não suportado',
                      response.get_data(as_text=True))

    def test_importar_conteudo_invalido(self):
        response = self._importar(BytesIO(b'lixo'), nome='planilha.xlsx')
        self.assertIn('Não foi possível ler o arquivo',
                      response.get_data(as_text=True))

    def test_importar_colunas_desconhecidas(self):
        wb = Workbook()
        wb.active.append(['A', 'B', 'C', 'D'])
        wb.active.append([1, 'Alguém', 'Sim', 2])
        saida = BytesIO()
        wb.save(saida)
        saida.seek(0)
        response = self._importar(saida)
        self.assertIn('não segue o modelo',
                      response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherNightShiftPay.query.count(), 0)

    def test_importar_sem_permissao_vira_403(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions
                                if p.code != 'an:create']
            db.session.commit()
        response = self._importar(self._planilha_bytes([]))
        self.assertEqual(response.status_code, 403)


if __name__ == '__main__':
    unittest.main()
