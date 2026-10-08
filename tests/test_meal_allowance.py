"""Testes do módulo Vale Alimentação - Professores (RH).

Lançamento simples (professor + dias trabalhados) na mesma página da
listagem, com modal de detalhes como o da Hora Extra. Regras cobertas:
validação dos campos, escopo por unidade, permissões (meal:read/meal:create)
e o toggle Financeiro da unidade.
"""
import os
import tempfile
import unittest
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook

from app import create_app
from app.config import Config
from app.extensions import db
from app.models import (Classroom, Permission, Reservation, Role,
                        RoomCategory, TeacherMealAllowance, Unity, User)

EMAIL = 'gestor@escola.edu'
PASSWORD = 'SenhaForte123'


class TestConfig(Config):
    SECRET_KEY = 'chave-de-teste-nao-usar-o-valor-dev'
    TESTING = True
    WTF_CSRF_ENABLED = False
    RATELIMIT_ENABLED = False


class MealAllowanceTestCase(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix='.db')
        os.close(fd)
        TestConfig.SQLALCHEMY_DATABASE_URI = 'sqlite:///' + self.db_path.replace('\\', '/')

        self.app = create_app(TestConfig)
        self.app.config['SESSION_COOKIE_SECURE'] = False
        self.client = self.app.test_client()

        with self.app.app_context():
            db.create_all()
            self.unity = Unity(name='Unidade Teste', code='UT')
            db.session.add(self.unity)
            db.session.flush()

            perms = [Permission(code='meal:read', module='meal', action='read'),
                     Permission(code='meal:create', module='meal', action='create'),
                     Permission(code='meal:edit', module='meal', action='edit'),
                     Permission(code='meal:delete', module='meal', action='delete')]
            db.session.add_all(perms)
            role = Role(name='gestor-teste', label='Gestor Teste', permissions=perms)
            db.session.add(role)
            db.session.flush()

            gestor = User(
                email=EMAIL, full_name='Gestor Teste',
                role='room', profile_type='employee', unities=[self.unity],
                role_id=role.id, force_password_change=False, is_active_user=True,
            )
            gestor.set_password(PASSWORD)
            db.session.add(gestor)
            db.session.flush()

            self.professor = User(
                email='professor@escola.edu', full_name='Professor Silva',
                role='room', profile_type='teacher', unities=[self.unity],
                force_password_change=False, is_active_user=True,
            )
            self.professor.set_password(PASSWORD)
            self.professor2 = User(
                email='professora@escola.edu', full_name='Professora Souza',
                role='room', profile_type='teacher', unities=[self.unity],
                force_password_change=False, is_active_user=True,
            )
            self.professor2.set_password(PASSWORD)
            db.session.add_all([self.professor, self.professor2])
            db.session.flush()

            # Sala para as reservas usadas nos testes de importação.
            categoria = RoomCategory(name='Sala de Aula', code='SA')
            db.session.add(categoria)
            db.session.flush()
            sala = Classroom(name='Sala 101', code='S101', capacity=30,
                             category_id=categoria.id,
                             unity_id=self.unity.id, is_active=True)
            db.session.add(sala)
            db.session.commit()
            self.gestor_id = gestor.id
            self.sala_id = sala.id
            self.unity_id = self.unity.id
            self.professor_id = self.professor.id
            self.professor2_id = self.professor2.id

        response = self.client.post('/login', data={'email': EMAIL, 'password': PASSWORD},
                                    follow_redirects=True)
        self.assertEqual(response.status_code, 200)

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        for suffix in ('', '-wal', '-shm'):
            path = self.db_path + suffix
            if os.path.exists(path):
                os.remove(path)

    def _adicionar(self, teacher=None, days='12', mes=None):
        data = {'teacher': str(self.professor_id if teacher is None
                               else teacher),
                'days': days}
        if mes is not None:
            data['month_base'] = mes
        return self.client.post('/payments/meal-allowance/add',
                                data=data, follow_redirects=True)

    def _id_unico(self):
        with self.app.app_context():
            return TeacherMealAllowance.query.one().id

    def _reserva(self, teacher_id, dia, mes=9, status='approved'):
        return Reservation(
            user_id=self.gestor_id, classroom_id=self.sala_id,
            title=f'Aula {teacher_id}-{dia}', date=date(2026, mes, dia),
            start_time=time(8, 0), end_time=time(10, 0),
            status=status, unity_id=self.unity_id, teacher_id=teacher_id)

    def _importar(self, month='2026-09'):
        return self.client.post('/payments/meal-allowance/import',
                                data={'month': month}, follow_redirects=True)

    def test_pagina_mostra_form_listagem_e_menu(self):
        response = self.client.get('/payments/meal-allowance')
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('Vale Alimentação', html)
        self.assertIn('Dias Trabalhados', html)
        self.assertIn('Professor Silva', html)  # choices do dropdown
        self.assertIn('Professora Souza', html)
        self.assertIn('Nenhum lançamento adicionado', html)
        # item do menu RH no sidebar
        self.assertIn('/payments/meal-allowance', html)
        # ações do cabeçalho: exportação e limpeza geral
        self.assertIn('/payments/meal-allowance/export', html)
        self.assertIn('/payments/meal-allowance/clear', html)

    def test_adicionar_lancamento(self):
        response = self._adicionar(days='12')
        self.assertIn('Lançamento adicionado', response.get_data(as_text=True))
        with self.app.app_context():
            lancamento = TeacherMealAllowance.query.one()
            self.assertEqual(lancamento.teacher_id, self.professor_id)
            self.assertEqual(lancamento.days, 12)
            self.assertEqual(lancamento.unity_id, self.unity_id)
            self.assertIsNotNone(lancamento.created_by_id)
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('Professor Silva', html)
        self.assertIn('12 dias', html)
        # Modal de detalhes e botões de ação, como na Hora Extra
        self.assertIn('view-meal-', html)
        self.assertIn('Responsável pelo Lançamento', html)
        self.assertIn('Excluir este lançamento', html)
        self.assertIn(f'/payments/meal-allowance/{self._id_unico()}/edit', html)

    def test_editar_lancamento(self):
        self._adicionar(days='12')
        entry_id = self._id_unico()

        response = self.client.get(f'/payments/meal-allowance/{entry_id}/edit')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Editar Lançamento', response.get_data(as_text=True))
        self.assertIn('value="12"', response.get_data(as_text=True))
        self.assertIn('Mês de Referência', response.get_data(as_text=True))

        response = self.client.post(f'/payments/meal-allowance/{entry_id}/edit',
                                    data={'teacher': str(self.professor2_id),
                                          'month_base': '2026-11',
                                          'days': '20'},
                                    follow_redirects=True)
        self.assertIn('Alteração realizada', response.get_data(as_text=True))
        with self.app.app_context():
            lancamento = db.session.get(TeacherMealAllowance, entry_id)
            self.assertEqual((lancamento.teacher_id, lancamento.days,
                              lancamento.month_base),
                             (self.professor2_id, 20, '2026-11'))

    def test_excluir_lancamento(self):
        self._adicionar(days='12')
        entry_id = self._id_unico()
        response = self.client.post(f'/payments/meal-allowance/{entry_id}/delete',
                                    follow_redirects=True)
        self.assertIn('excluído', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_filtro_por_professor(self):
        self._adicionar(teacher=self.professor_id, days='12')
        self._adicionar(teacher=self.professor2_id, days='5')

        sem_filtro = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('12 dias', sem_filtro)
        self.assertIn('5 dias', sem_filtro)

        filtrado = self.client.get(
            f'/payments/meal-allowance?teacher_filter={self.professor2_id}'
        ).get_data(as_text=True)
        self.assertIn('5 dias', filtrado)
        self.assertNotIn('12 dias', filtrado)
        # O select do filtro mantém a seleção aplicada
        self.assertIn(f'<option value="{self.professor2_id}" selected>', filtrado)

    def test_dias_invalidos_recusados(self):
        # Zero tem input: cai na validação de faixa, não na de campo vazio.
        response = self._adicionar(days='0')
        self.assertIn('entre 1 e 999', response.get_data(as_text=True))
        response = self._adicionar(days='1000')
        self.assertIn('entre 1 e 999', response.get_data(as_text=True))
        response = self._adicionar(days='')
        self.assertIn('Informe os dias trabalhados', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_professor_obrigatorio(self):
        response = self._adicionar(teacher=0)
        self.assertIn('Selecione o professor', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_escopo_unidade(self):
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            db.session.add(TeacherMealAllowance(
                teacher_id=self.professor_id, days=5, unity_id=outra.id))
            db.session.commit()
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('Nenhum lançamento adicionado', html)

    def test_sem_permissao_vira_403(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = []
            db.session.commit()
        self.assertEqual(self.client.get('/payments/meal-allowance').status_code, 403)
        self.assertEqual(self._adicionar().status_code, 403)

    def test_leitura_sem_criacao_esconde_form(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions if p.code == 'meal:read']
            db.session.commit()
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertNotIn('payments.add_meal_allowance', html)
        self.assertIn('Nenhum lançamento adicionado', html)
        self.assertEqual(self._adicionar().status_code, 403)

    def test_editar_excluir_sem_permissao_vira_403(self):
        self._adicionar(days='12')
        entry_id = self._id_unico()
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions
                                if p.code in ('meal:read', 'meal:create')]
            db.session.commit()
        self.assertEqual(
            self.client.get(f'/payments/meal-allowance/{entry_id}/edit').status_code, 403)
        self.assertEqual(
            self.client.post(f'/payments/meal-allowance/{entry_id}/edit',
                             data={'teacher': str(self.professor_id), 'days': '20'}).status_code,
            403)
        self.assertEqual(
            self.client.post(f'/payments/meal-allowance/{entry_id}/delete').status_code,
            403)
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 1)

    def test_modulo_financeiro_desligado_vira_403(self):
        with self.app.app_context():
            unity = db.session.get(Unity, self.unity_id)
            unity.finance_enabled = False
            db.session.commit()
        self.assertEqual(self.client.get('/payments/meal-allowance').status_code, 403)
        self.assertEqual(self._adicionar().status_code, 403)

    # ---------- Exportação Excel ----------

    def _exportar(self, query='', **kwargs):
        return self.client.get('/payments/meal-allowance/export' + query, **kwargs)

    def _linhas_planilha(self, response):
        wb = load_workbook(BytesIO(response.data))
        return list(wb.active.iter_rows(values_only=True))

    def test_exportar_excel_total_por_professor(self):
        self._adicionar(teacher=self.professor_id, days='12')
        self._adicionar(teacher=self.professor_id, days='3')
        self._adicionar(teacher=self.professor2_id, days='5')

        response = self._exportar()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('vale_alimentacao_professores',
                      response.headers['Content-Disposition'])

        rows = self._linhas_planilha(response)
        self.assertEqual(rows[0], ('Professor', 'Dias Trabalhados'))
        dados = {nome: dias for nome, dias in rows[1:]}
        # Lançamentos do mesmo professor são somados numa linha só.
        self.assertEqual(dados['Professor Silva'], 15)
        self.assertEqual(dados['Professora Souza'], 5)
        self.assertEqual(len(rows), 3)

    def test_exportar_respeita_filtro_por_professor(self):
        self._adicionar(teacher=self.professor_id, days='12')
        self._adicionar(teacher=self.professor2_id, days='5')

        response = self._exportar(f'?teacher_filter={self.professor2_id}')
        rows = self._linhas_planilha(response)
        self.assertEqual(rows[1], ('Professora Souza', 5))
        self.assertEqual(len(rows), 2)

    def test_exportar_escopo_unidade(self):
        self._adicionar(days='12')
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            db.session.add(TeacherMealAllowance(
                teacher_id=self.professor_id, days=5, unity_id=outra.id))
            db.session.commit()

        rows = self._linhas_planilha(self._exportar())
        self.assertEqual(rows[1], ('Professor Silva', 12))
        self.assertEqual(len(rows), 2)

    def test_exportar_sem_lancamentos_redireciona(self):
        response = self._exportar()
        self.assertEqual(response.status_code, 302)
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('Nenhum lançamento de Vale Alimentação para exportar', html)

    def test_exportar_sem_permissao_vira_403(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = []
            db.session.commit()
        self.assertEqual(self._exportar().status_code, 403)

    # ---------- Limpar todos os lançamentos ----------

    def test_limpar_todos_lancamentos(self):
        self._adicionar(teacher=self.professor_id, days='12')
        self._adicionar(teacher=self.professor2_id, days='5')

        response = self.client.post('/payments/meal-allowance/clear',
                                    follow_redirects=True)
        self.assertIn('2 lançamento(s) do Vale Alimentação removido(s)',
                      response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_limpar_todos_escopo_unidade(self):
        self._adicionar(days='12')
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            db.session.add(TeacherMealAllowance(
                teacher_id=self.professor_id, days=5, unity_id=outra.id))
            db.session.commit()
            outra_id = outra.id

        self.client.post('/payments/meal-allowance/clear')
        with self.app.app_context():
            # Só o lançamento de outra unidade sobrevive.
            restante = TeacherMealAllowance.query.one()
            self.assertEqual(restante.unity_id, outra_id)

    def test_limpar_todos_sem_permissao_vira_403(self):
        self._adicionar(days='12')
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions if p.code == 'meal:read']
            db.session.commit()

        # meal:read sozinho vê a página, mas não tem o botão nem a rota.
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertNotIn('/payments/meal-allowance/clear', html)
        self.assertEqual(
            self.client.post('/payments/meal-allowance/clear').status_code, 403)
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 1)

    def test_cabecalho_com_lancamentos(self):
        """Com lançamentos, a confirmação do Limpar Tudo mostra a contagem
        total da unidade e a exportação herda o filtro de professor ativo."""
        self._adicionar(teacher=self.professor_id, days='12')
        self._adicionar(teacher=self.professor2_id, days='5')

        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('Excluir TODOS os 2 lançamento(s)', html)

        filtrado = self.client.get(
            f'/payments/meal-allowance?teacher_filter={self.professor2_id}'
        ).get_data(as_text=True)
        self.assertIn(
            f'/payments/meal-allowance/export?teacher_filter={self.professor2_id}',
            filtrado)

    # ---------- Importar dos Agendamentos ----------

    def test_caixa_importacao_na_pagina(self):
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('Listar Professores', html)
        self.assertIn('/payments/meal-allowance/import', html)
        self.assertIn('Mês dos Agendamentos', html)
        # Mês corrente vem selecionado por padrão mesmo sem reservas.
        agora = datetime.now()
        self.assertIn(f'value="{agora:%Y-%m}" selected', html)

    def test_caixa_importacao_sem_permissao_de_criacao(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions if p.code == 'meal:read']
            db.session.commit()
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertNotIn('/payments/meal-allowance/import', html)

    def test_importar_professores_dos_agendamentos(self):
        with self.app.app_context():
            db.session.add_all([
                # Silva: 3 datas distintas no mês + 1 cancelada (ignorada).
                self._reserva(self.professor_id, 10),
                self._reserva(self.professor_id, 11),
                self._reserva(self.professor_id, 12),
                self._reserva(self.professor_id, 13, status='cancelled'),
                # Souza: 2 reservas no mesmo dia contam 1 dia + 1 pendente (ignorada).
                self._reserva(self.professor2_id, 10),
                self._reserva(self.professor2_id, 10),
                self._reserva(self.professor2_id, 14, status='pending'),
            ])
            db.session.commit()

        response = self._importar()
        html = response.get_data(as_text=True)
        self.assertIn('2 lançamento(s) criado(s) a partir dos agendamentos de '
                      'Setembro/2026', html)
        self.assertIn('Professor Silva', html)
        self.assertIn('3 dias', html)
        self.assertIn('Professora Souza', html)
        self.assertIn('1 dia<', html)  # singular, sem 's'
        with self.app.app_context():
            lancamentos = {l.teacher_id: l.days
                           for l in TeacherMealAllowance.query.all()}
            self.assertEqual(lancamentos[self.professor_id], 3)
            self.assertEqual(lancamentos[self.professor2_id], 1)

    def test_importar_nao_duplica_lancamentos(self):
        with self.app.app_context():
            db.session.add(self._reserva(self.professor_id, 10))
            db.session.commit()
        self._importar()
        response = self._importar()
        html = response.get_data(as_text=True)
        self.assertIn('já possuíam lançamento no mês e foram mantidos: Professor Silva',
                      html)
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 1)

    def test_importar_grava_mes_e_aceita_outro_mes(self):
        """O lançamento fica no mês importado; importar outro mês cria
        lançamento novo para o mesmo professor, sem duplicar o primeiro."""
        with self.app.app_context():
            db.session.add_all([
                self._reserva(self.professor_id, 10, mes=9),
                self._reserva(self.professor_id, 15, mes=10),
            ])
            db.session.commit()
        self._importar(month='2026-09')
        self._importar(month='2026-10')
        with self.app.app_context():
            por_mes = {l.month_base: l.days
                       for l in TeacherMealAllowance.query.all()}
            self.assertEqual(por_mes, {'2026-09': 1, '2026-10': 1})

    def test_importar_sem_agendamentos_no_mes(self):
        response = self._importar(month='2026-01')
        html = response.get_data(as_text=True)
        self.assertIn('Nenhum professor com agendamento aprovado em '
                      'Janeiro/2026', html)
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_importar_escopo_unidade(self):
        with self.app.app_context():
            outra = Unity(name='Outra Unidade', code='OU')
            db.session.add(outra)
            db.session.flush()
            categoria = RoomCategory.query.first()
            sala = Classroom(name='Sala Norte', code='N201', capacity=20,
                             category_id=categoria.id, unity_id=outra.id,
                             is_active=True)
            db.session.add(sala)
            db.session.flush()
            db.session.add(Reservation(
                user_id=self.gestor_id, classroom_id=sala.id,
                title='Aula em outra unidade', date=date(2026, 9, 10),
                start_time=time(8, 0), end_time=time(10, 0),
                status='approved', unity_id=outra.id,
                teacher_id=self.professor_id))
            db.session.commit()

        response = self._importar()
        self.assertIn('Nenhum professor com agendamento aprovado',
                      response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_importar_mes_invalido(self):
        response = self._importar(month='banana')
        self.assertIn('Selecione um mês válido',
                      response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.count(), 0)

    def test_importar_sem_permissao_vira_403(self):
        with self.app.app_context():
            role = Role.query.filter_by(name='gestor-teste').first()
            role.permissions = [p for p in role.permissions if p.code == 'meal:read']
            db.session.commit()
        self.assertEqual(self._importar().status_code, 403)

    def test_importar_modulo_financeiro_desligado_vira_403(self):
        with self.app.app_context():
            unity = db.session.get(Unity, self.unity_id)
            unity.finance_enabled = False
            db.session.commit()
        self.assertEqual(self._importar().status_code, 403)

    def test_seletor_lista_meses_com_reservas(self):
        with self.app.app_context():
            db.session.add(self._reserva(self.professor_id, 10, mes=3))
            db.session.commit()
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('value="2026-03"', html)
        self.assertIn('Março/2026', html)

    # ---------- Mês de referência ----------

    def test_adicionar_com_mes_de_referencia(self):
        response = self._adicionar(days='12', mes='2026-08')
        html = response.get_data(as_text=True)
        self.assertIn('12 dia(s) trabalhados em Agosto/2026', html)
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.one().month_base,
                             '2026-08')
        self.assertIn('Agosto/2026',
                      self.client.get('/payments/meal-allowance').get_data(as_text=True))

    def test_adicionar_sem_mes_usa_mes_corrente(self):
        agora = datetime.now()
        self._adicionar(days='5')
        with self.app.app_context():
            self.assertEqual(TeacherMealAllowance.query.one().month_base,
                             f'{agora:%Y-%m}')

    def test_lancamento_legado_sem_mes_exibe_traco(self):
        """Lançamentos anteriores à coluna (month_base nulo) continuam na
        listagem, com o mês exibido como '—'."""
        with self.app.app_context():
            db.session.add(TeacherMealAllowance(
                teacher_id=self.professor_id, days=4,
                unity_id=self.unity_id, created_by_id=self.gestor_id))
            db.session.commit()
        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('4 dias', html)
        self.assertIn('<td>—</td>', html)

    def test_filtro_por_mes(self):
        self._adicionar(teacher=self.professor_id, days='12', mes='2026-08')
        self._adicionar(teacher=self.professor_id, days='5', mes='2026-09')

        html = self.client.get('/payments/meal-allowance').get_data(as_text=True)
        self.assertIn('12 dias', html)
        self.assertIn('5 dias', html)
        self.assertIn('Todos os Meses', html)

        filtrado = self.client.get(
            '/payments/meal-allowance?month_base=2026-08').get_data(as_text=True)
        self.assertIn('12 dias', filtrado)
        self.assertNotIn('5 dias', filtrado)
        self.assertIn('<option value="2026-08" selected>', filtrado)

    def test_exportar_respeita_filtro_por_mes(self):
        self._adicionar(teacher=self.professor_id, days='12', mes='2026-08')
        self._adicionar(teacher=self.professor_id, days='5', mes='2026-09')

        rows = self._linhas_planilha(self._exportar('?month_base=2026-09'))
        self.assertEqual(rows[1], ('Professor Silva', 5))
        self.assertEqual(len(rows), 2)


if __name__ == '__main__':
    unittest.main()
