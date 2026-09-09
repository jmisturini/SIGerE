"""Testes do recálculo de quantidades persistido no módulo Cozinha.

Cobre o botão "Salvar quantidades" da preparação: a escala escolhida fica
salva na receita, a exibição e a requisição de compras passam a usar as
quantidades recalculadas, e "Restaurar originais" volta ao estado da ficha —
sem sobrescrever os ingredientes originais.
"""
import os
import tempfile
import unittest
from io import BytesIO

from openpyxl import load_workbook

from app import create_app
from app.config import Config
from app.extensions import db
from app.models import (KitchenPreparation, KitchenRecipe, KitchenRecipeIngredient,
                        Permission, Role, Unity, User)

USERNAME = 'cozinha.teste'
PASSWORD = 'SenhaForte123'


class TestConfig(Config):
    SECRET_KEY = 'chave-de-teste-nao-usar-o-valor-dev'
    TESTING = True
    WTF_CSRF_ENABLED = False
    RATELIMIT_ENABLED = False


class KitchenScaleTestCase(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix='.db')
        os.close(fd)
        TestConfig.SQLALCHEMY_DATABASE_URI = 'sqlite:///' + self.db_path.replace('\\', '/')

        self.app = create_app(TestConfig)
        self.app.config['SESSION_COOKIE_SECURE'] = False
        self.client = self.app.test_client()

        with self.app.app_context():
            db.create_all()
            self.unity = Unity(name='Unidade Teste', code='UT')
            db.session.add(self.unity)
            db.session.flush()

            perms = [Permission(code=code, module='kitchen', action=code.split(':')[1])
                     for code in ('kitchen:read', 'kitchen:sheet_create',
                                  'kitchen:shopping_export')]
            db.session.add_all(perms)
            role = Role(name='cozinheiro-teste', label='Cozinha Teste', permissions=perms)
            db.session.add(role)
            db.session.flush()

            user = User(
                username=USERNAME, email='cozinha@escola.edu', full_name='Cozinheiro Teste',
                role='room', profile_type='employee', unity_id=self.unity.id,
                role_id=role.id, force_password_change=False, is_active_user=True,
            )
            user.set_password(PASSWORD)
            db.session.add(user)

            # Receita com rendimento em faixa: base = 4 porções
            recipe = KitchenRecipe(name='Bolo de Carne', unity_id=self.unity.id,
                                   yield_info='4 a 6 porções')
            preparation = KitchenPreparation(name='Bolo de Carne', position=0, recipe=recipe)
            preparation.ingredients.append(KitchenRecipeIngredient(
                name='Farinha de Trigo', quantity=400.0, unit='g', position=0))
            preparation.ingredients.append(KitchenRecipeIngredient(
                name='Sal', quantity=None, quantity_raw='a gosto', unit='', position=1))
            db.session.add(recipe)
            db.session.commit()
            self.recipe_id = recipe.id

        self._login()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        for suffix in ('', '-wal', '-shm'):
            path = self.db_path + suffix
            if os.path.exists(path):
                os.remove(path)

    def _login(self):
        response = self.client.post('/login', data={'username': USERNAME, 'password': PASSWORD},
                                    follow_redirects=True)
        self.assertEqual(response.status_code, 200)

    def _get_recipe(self):
        with self.app.app_context():
            recipe = db.session.get(KitchenRecipe, self.recipe_id)
            db.session.refresh(recipe)
            return recipe

    def test_save_portions_persists_scale(self):
        response = self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                                    data={'portions': '30'}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn('Quantidades recalculadas para 30 porções e salvas', page)
        recipe = self._get_recipe()
        self.assertEqual(recipe.scaled_portions, 30.0)
        self.assertAlmostEqual(recipe.scale_factor, 7.5)
        # A página passa a exibir as quantidades escaladas e o selo de salvo
        self.assertIn('3000', page)          # 400 g × 7,5
        self.assertIn('Salvo: 30 porções', page)
        self.assertIn('Restaurar originais', page)

    def test_detail_shows_original_without_scale(self):
        response = self.client.get(f'/kitchen/preparacoes/{self.recipe_id}')
        page = response.get_data(as_text=True)
        self.assertIn('>400<', page)
        self.assertNotIn('Salvo:', page)

    def test_save_with_base_portions_clears_scale(self):
        self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                         data={'portions': '30'})
        # Salvar o próprio rendimento base equivale a não ter escala
        response = self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                                    data={'portions': '4'}, follow_redirects=True)
        self.assertIn('restauradas ao rendimento original', response.get_data(as_text=True))
        recipe = self._get_recipe()
        self.assertIsNone(recipe.scaled_portions)

    def test_restore_portions_clears_scale(self):
        self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                         data={'portions': '30'})
        response = self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/restaurar',
                                    follow_redirects=True)
        self.assertIn('Quantidades originais da ficha restauradas', response.get_data(as_text=True))
        recipe = self._get_recipe()
        self.assertIsNone(recipe.scaled_portions)

    def test_shopping_export_uses_scaled_quantities(self):
        from app.blueprints.kitchen.export import aggregate_ingredients

        self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                         data={'portions': '30'})
        with self.app.app_context():
            recipe = db.session.get(KitchenRecipe, self.recipe_id)
            rows = aggregate_ingredients([recipe])
            farinha = next(r for r in rows if r['nome'] == 'Farinha de Trigo')
            self.assertAlmostEqual(farinha['quantidade'], 3.0)   # 400 g × 7,5 = 3000 g = 3 KG
            self.assertEqual(farinha['unidade'], 'KG')

    def test_shopping_page_shows_scale_badge(self):
        self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                         data={'portions': '30'})
        response = self.client.get('/kitchen/compras')
        self.assertIn('quantidades para 30 porções', response.get_data(as_text=True))

    def test_exported_xlsx_has_scaled_quantity(self):
        self.client.post(f'/kitchen/preparacoes/{self.recipe_id}/porcoes/salvar',
                         data={'portions': '30'})
        response = self.client.post('/kitchen/compras/export',
                                    data={'recipe_ids': str(self.recipe_id),
                                          'professor': 'Cozinheiro Teste',
                                          'class_date': '2026-09-10'})
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data))
        ws = workbook[workbook.sheetnames[0]]
        quantities = [ws.cell(row=row, column=2).value for row in range(10, ws.max_row + 1)]
        self.assertIn(3.0, quantities)   # 400 g × 7,5 → 3 KG


if __name__ == '__main__':
    unittest.main()
