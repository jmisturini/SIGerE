"""Módulo Adicional Noturno de Professores (RH).

Adaptação web do gerador local "Gerador-Ad-Noturno-SENAC-FL"
(Adicional Noturno.py): o formulário replica as colunas da planilha
"Adicional Noturno Professores" (Matrícula, Professor, Adicional Noturno
Sim/Não, Horas/Aula, Minutos e Unidade) e a exportação preenche o modelo
app/static/templates_excel/planilha_base_an.xlsx ("TABELA PARA PAGAMENTO
ADICIONAL NOTURNO PROFESSOR") nos mesmos moldes do script original —
Matrícula, Nome, H/A e Minutos Decimais na coluna G e UO Lotação na
coluna K, a partir da linha 5. Só entra na planilha quem tem adicional
"Sim" com minutos maiores que zero, ordenado pelo código da unidade; as
demais colunas seguem com as fórmulas do modelo (dias, total h/m e valor
do adicional).
"""
import io
import os
import re
import unicodedata
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree as ET

from flask import (Blueprint, abort, current_app, flash, redirect,
                   render_template, request, send_file, url_for)
from openpyxl import load_workbook
from sqlalchemy import func

from app.extensions import db
from flask_login import current_user, login_required
from app.forms import FormAdicionalNoturno
from app.models import TeacherNightShiftPay, User
from app.permissions import require_module, require_permission
from app.blueprints.payments import (_mes_anterior_de, _mes_atual,
                                     _mes_valido, _proximo_mes, _rotulo_mes,
                                     _teachers_for_current_unity)
from app.unity_context import current_unity_id
from app.utils import redirect_back, redirect_preserving_args

bp = Blueprint('an', __name__, url_prefix='/adicional-noturno')

# Modelo da planilha de pagamento: aba "Adicional Noturno", dados da linha 5.
AN_SOURCE_SHEET = 'Adicional Noturno'
AN_HIDDEN_SHEET = 'Não mexer nesta planilha'
AN_TEMPLATE_FILE = 'planilha_base_an.xlsx'
EXPORT_FIRST_ROW = 5
# As fórmulas do modelo vão até a linha 100 e a aba oculta "Não mexer nesta
# planilha" referencia exatamente esse intervalo — acima disso o arquivo
# gerado perde o vínculo (mesmo cuidado do modelo do Vale-Transporte).
EXPORT_MAX_ROWS = 96

# Namespaces do OOXML usados na reposição do cache das fórmulas.
_SSML = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
_REL_DOC = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_REL_PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'

# Códigos UO do Senac SC — espelho das colunas M/N da aba "Adicional
# Noturno" do modelo, oferecido como referência no formulário (a planilha
# de origem usa o número, sem o nome da unidade).
UO_CODES = (
    (2, 'Florianópolis'), (3, 'Joinville'), (4, 'Blumenau'), (5, 'Lages'),
    (6, 'Chapecó'), (7, 'Brusque'), (8, 'Caçador'), (9, 'Criciúma'),
    (10, 'Concórdia'), (11, 'Itajaí'), (12, 'Jaraguá do Sul'),
    (13, 'Joaçaba'), (14, 'São Miguel do Oeste'), (15, 'Tubarão'),
    (16, 'Rio do Sul'), (17, 'São Bento do Sul'), (18, 'Xanxerê'),
    (19, 'Canoinhas'), (21, 'TI'), (22, 'Bistrô'), (23, 'Saúde e Beleza'),
    (29, 'DEP'), (33, 'EAD'), (34, 'Curitibanos'), (35, 'Porto União'),
    (36, 'Palhoça'), (37, 'Timbó'), (38, 'Videira'), (39, 'Mafra'),
    (40, 'DIT'), (41, 'Araranguá'), (42, 'Campos Novos'),
)

REGISTROS_PER_PAGE = 25


# ── Helpers ──────────────────────────────────────────────────────────────────

def _meses_janela():
    """Janela fixa dos seletores: os 12 meses anteriores, o mês corrente e o
    seguinte — garante meses vizinhos selecionáveis ao corrigir um lançamento."""
    meses = [_proximo_mes()]
    mes = _mes_atual()
    for _ in range(13):
        meses.append(mes)
        mes = _mes_anterior_de(mes)
    return meses


def _meses_opcoes(unity_id):
    """Opções do seletor de mês do formulário e do filtro: a janela fixa
    (12 anteriores + corrente + seguinte) mais os meses que já têm
    lançamentos na unidade, do mais recente para o mais antigo."""
    meses = set(_meses_janela())
    lancamentos = (db.session.query(TeacherNightShiftPay.month_base)
                   .filter(TeacherNightShiftPay.unity_id == unity_id)
                   .distinct().all())
    meses.update(row[0] for row in lancamentos)
    return [(valor, _rotulo_mes(valor)) for valor in sorted(meses, reverse=True)]


def _get_registro_scoped(registro_id):
    """Carrega o lançamento da unidade ativa — de outras unidades dá 404."""
    registro = db.get_or_404(TeacherNightShiftPay, registro_id)
    if registro.unity_id != current_unity_id():
        abort(404)
    return registro


# ── Seleção de professor do cadastro ─────────────────────────────────────────

MANUAL = 'manual'  # opção "Outro" do select: digitação livre de nome/matrícula

# Choices iniciais do select; os professores entram entre elas na rota.
PROFESSOR_CHOICES_BASE = [('', 'Selecione…'),
                          (MANUAL, 'Outro (digitar nome e matrícula)')]


def _professores_choices():
    """Professores ativos da unidade ativa para o select, em ordem alfabética."""
    return [(str(u.id), u.full_name) for u in _teachers_for_current_unity()]


def _professores_map():
    """{id: {nome, matricula}} dos professores da unidade ativa — alimenta o
    auto-preenchimento da matrícula no JavaScript da página."""
    return {str(u.id): {'nome': u.full_name, 'matricula': u.registration or ''}
            for u in _teachers_for_current_unity()}


def _professor_id_por_matricula(matricula):
    """Id (como texto) do professor da unidade ativa com a matrícula dada —
    para pré-selecionar o select na edição. None quando ninguém tem."""
    for u in _teachers_for_current_unity():
        if u.registration == matricula:
            return str(u.id)
    return None


def _resolver_professor(form):
    """(nome, matrícula) do lançamento: professor selecionado do cadastro —
    nome e matrícula vêm da CONTA, não do formulário — ou a opção "Outro",
    com digitação manual dos dois. Levanta ValueError com a mensagem de
    validação quando a escolha não fecha."""
    matricula = (form.registration.data or '').strip()
    if form.teacher.data == MANUAL:
        nome = (form.full_name.data or '').strip()
        if not nome:
            raise ValueError('Informe o nome do professor (ou selecione um '
                             'professor do cadastro).')
        if not matricula:
            raise ValueError('Informe a matrícula.')
        return nome, matricula
    if not form.teacher.data:
        raise ValueError('Selecione o professor.')
    professor = db.session.get(User, int(form.teacher.data))
    if professor is None:
        raise ValueError('Selecione o professor.')
    # A matrícula do cadastro é a autoritativa; o campo digitado vale apenas
    # para professor sem matrícula gravada (o auto-preenchimento o libera).
    if professor.registration:
        matricula = professor.registration
    if not matricula:
        raise ValueError(f'{professor.full_name} não tem matrícula no '
                         'cadastro — informe a matrícula.')
    return professor.full_name, matricula


def _class_hours_decimal(texto):
    """'58,91' → Decimal('58.91'); vazio/inválido vira None (o modelo trata
    a coluna C vazia como zero nas fórmulas, como na planilha original)."""
    try:
        return Decimal(texto.strip().replace(',', '.')) if texto else None
    except InvalidOperation:
        return None


def _preparar_form(form):
    """Choices dos seletores de mês e professor a partir da unidade ativa."""
    form.month_base.choices = _meses_opcoes(current_unity_id())
    form.teacher.choices = (PROFESSOR_CHOICES_BASE[:1]
                            + _professores_choices()
                            + PROFESSOR_CHOICES_BASE[1:])
    return form


def _render_registros(form, month_filter=None):
    """Página completa do módulo (formulário + filtro + listagem) — usada no
    GET e no re-render do POST com erro de validação."""
    query = TeacherNightShiftPay.query.filter_by(unity_id=current_unity_id())
    if month_filter:
        query = query.filter_by(month_base=month_filter)
    # Ordenação escolhida no filtro da página (padrão: lançamento recente);
    # a matrícula é texto — ordena por tamanho antes do valor (20 antes de 100).
    ordem = request.args.get('ordem', 'recentes')
    if ordem == 'antigas':
        query = query.order_by(TeacherNightShiftPay.created_at,
                               TeacherNightShiftPay.id)
    elif ordem == 'professor':
        query = query.order_by(TeacherNightShiftPay.full_name,
                               TeacherNightShiftPay.month_base)
    elif ordem == 'matricula':
        query = query.order_by(func.length(TeacherNightShiftPay.registration),
                               TeacherNightShiftPay.registration,
                               TeacherNightShiftPay.full_name)
    elif ordem == 'mes':
        query = query.order_by(TeacherNightShiftPay.month_base,
                               TeacherNightShiftPay.full_name)
    elif ordem == 'mes_desc':
        query = query.order_by(TeacherNightShiftPay.month_base.desc(),
                               TeacherNightShiftPay.full_name)
    else:
        ordem = 'recentes'
        query = query.order_by(TeacherNightShiftPay.created_at.desc(),
                               TeacherNightShiftPay.id.desc())
    pagination = query \
        .paginate(page=request.args.get('page', 1, type=int),
                  per_page=REGISTROS_PER_PAGE, error_out=False)
    # Total da unidade (sem filtro): alimenta a confirmação do "Limpar Tudo".
    total_registros = TeacherNightShiftPay.query \
        .filter_by(unity_id=current_unity_id()).count()
    return render_template('an/registros.html', form=form,
                           registros=pagination.items, pagination=pagination,
                           filter_month=month_filter,
                           ordem=ordem,
                           meses_opcoes=_meses_opcoes(current_unity_id()),
                           mes_corrente=_mes_atual(),
                           total_registros=total_registros,
                           uo_codes=UO_CODES,
                           professores_map=_professores_map())


# ── Cache dos resultados das fórmulas do modelo ──────────────────────────────

def _numero_estatico(celulas, ref):
    """Valor numérico já salvo na célula (estático, sem fórmula); 0.0 quando
    a célula está em branco ou é uma fórmula."""
    celula = celulas.get(ref)
    if celula is None or celula.find(f'{{{_SSML}}}f') is not None:
        return 0.0
    v = celula.find(f'{{{_SSML}}}v')
    try:
        return float(v.text) if v is not None and v.text else 0.0
    except (TypeError, ValueError):
        return 0.0


def _valor_estatico(celulas, ref):
    """Valor salvo na célula (número ou texto); None quando em branco. Cobre
    os formatos que o openpyxl grava: número/texto simples em <v> (t='str')
    e texto embutido em <is><t> (t='inlineStr')."""
    celula = celulas.get(ref)
    if celula is None or celula.find(f'{{{_SSML}}}f') is not None:
        return None
    if celula.get('t') == 'inlineStr':
        embutida = celula.find(f'{{{_SSML}}}is')
        t = (embutida.find(f'{{{_SSML}}}t') if embutida is not None else None)
        return (t.text or '') if t is not None else None
    v = celula.find(f'{{{_SSML}}}v')
    if v is None or v.text is None:
        return None
    if celula.get('t') == 'str':
        return v.text
    try:
        return float(v.text)
    except (TypeError, ValueError):
        return v.text


def _escrever_cache(celula, valor):
    """Grava o resultado calculado como cache da fórmula da célula, sem
    tocar na fórmula. O openpyxl deixa <v/> vazio nas células de fórmula —
    é exatamente esse cache que a exportação repõe. Texto recebe t='str'
    (como o Excel cacheia referências a matrícula); vazio vira 0 numérico,
    como o Excel devolve de uma referência a célula em branco."""
    if celula is None or celula.find(f'{{{_SSML}}}f') is None:
        return
    v = celula.find(f'{{{_SSML}}}v')
    if v is None:
        v = ET.SubElement(celula, f'{{{_SSML}}}v')
    if v.text:
        return  # já há resultado em cache
    if isinstance(valor, str):
        celula.set('t', 'str')
        v.text = valor
    else:
        v.text = f'{float(valor or 0):.10g}'


def _preencher_cache_formulas(conteudo):
    """Devolve o XLSX com os resultados em cache das fórmulas do modelo.

    O openpyxl preserva as fórmulas, mas descarta os resultados que o Excel
    guarda junto delas — sem esse cache, as colunas calculadas do modelo
    (20%, total horas, total h/m e TOTAL Adic Not) aparecem EM BRANCO até o
    Excel recalcular, o que não acontece no Modo de Exibição Protegido.
    Repõe o cache recalculando as fórmulas fixas do modelo na aba de dados
    (E=C*D, H=G/G2, I=H*F, J=E*I) e as referências da aba oculta às linhas
    da planilha (B=K, C=A, E=I, G=J), sem alterar nenhuma fórmula — o Excel
    segue recalculando normalmente quando o arquivo é editado."""
    with zipfile.ZipFile(io.BytesIO(conteudo)) as arquivo:
        entradas = {nome: arquivo.read(nome) for nome in arquivo.namelist()}

    pasta = ET.fromstring(entradas['xl/workbook.xml'])
    relacoes = ET.fromstring(entradas['xl/_rels/workbook.xml.rels'])
    alvo_por_rid = {rel.get('Id'): rel.get('Target')
                    for rel in relacoes.iter(f'{{{_REL_PKG}}}Relationship')}
    caminho = {}
    for aba in pasta.iter(f'{{{_SSML}}}sheet'):
        alvo = alvo_por_rid[aba.get(f'{{{_REL_DOC}}}id')]
        caminho[aba.get('name')] = (alvo.lstrip('/') if alvo.startswith('/')
                                    else 'xl/' + alvo)

    ET.register_namespace('', _SSML)
    dados = ET.fromstring(entradas[caminho[AN_SOURCE_SHEET]])
    celulas = {c.get('r'): c for c in dados.iter(f'{{{_SSML}}}c')}
    g2 = _numero_estatico(celulas, 'G2')

    # Aba de dados: recalcula as fórmulas de TODAS as linhas do modelo — as
    # linhas sem lançamento ficam com 0,00, igual à planilha original salva
    # pelo Excel.
    totais = {}
    for linha in range(EXPORT_FIRST_ROW, EXPORT_FIRST_ROW + EXPORT_MAX_ROWS):
        c = _numero_estatico(celulas, f'C{linha}')
        d = _numero_estatico(celulas, f'D{linha}')
        dias = _numero_estatico(celulas, f'F{linha}')
        g = _numero_estatico(celulas, f'G{linha}')
        e = c * d
        h = g / g2 if g2 else 0.0
        i = h * dias
        totais[linha] = (i, e * i)
        for ref, valor in ((f'E{linha}', e), (f'H{linha}', h),
                           (f'I{linha}', i), (f'J{linha}', e * i)):
            _escrever_cache(celulas.get(ref), valor)

    # Aba oculta: repõe o cache das referências às linhas da planilha.
    oculta = ET.fromstring(entradas[caminho[AN_HIDDEN_SHEET]])
    celulas_ocultas = {c.get('r'): c for c in oculta.iter(f'{{{_SSML}}}c')}
    for linha in range(1, EXPORT_MAX_ROWS + 1):
        linha_dados = EXPORT_FIRST_ROW + linha - 1
        i, j = totais[linha_dados]
        _escrever_cache(celulas_ocultas.get(f'B{linha}'),
                        _valor_estatico(celulas, f'K{linha_dados}'))
        _escrever_cache(celulas_ocultas.get(f'C{linha}'),
                        _valor_estatico(celulas, f'A{linha_dados}'))
        _escrever_cache(celulas_ocultas.get(f'E{linha}'), i)
        _escrever_cache(celulas_ocultas.get(f'G{linha}'), j)

    entradas[caminho[AN_SOURCE_SHEET]] = ET.tostring(
        dados, encoding='UTF-8', xml_declaration=True)
    entradas[caminho[AN_HIDDEN_SHEET]] = ET.tostring(
        oculta, encoding='UTF-8', xml_declaration=True)

    saida = io.BytesIO()
    with zipfile.ZipFile(saida, 'w', zipfile.ZIP_DEFLATED) as novo:
        for nome, conteudo_entrada in entradas.items():
            novo.writestr(nome, conteudo_entrada)
    return saida.getvalue()


def _ajustar_larguras(worksheet, larguras):
    """Alarga as colunas que receberam dados quando o conteúdo excede a
    largura do modelo — sem nunca estreitar, preservando o layout do modelo
    fornecido. O modelo foi feito para linhas sem valor (a coluna E, do 20%,
    tem 4,71 e exibe '###' quando o cálculo deixa de ser '0,00')."""
    for letra, texto in larguras.items():
        atual = worksheet.column_dimensions[letra].width or 8.43
        necessario = texto + 2.5
        if necessario > atual:
            worksheet.column_dimensions[letra].width = round(necessario, 2)


# ── Importação da planilha "Adicional Noturno Professores" ───────────────────

def _normalizar(texto):
    """'Adicional Noturno (Minutos)' → 'adicional noturno (minutos)':
    minúsculas, sem acentos, espaços colapsados."""
    texto = ''.join(c for c in unicodedata.normalize('NFD', str(texto or ''))
                    if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', texto).strip().lower()


def _colunas_do_modelo(aba):
    """Localiza as colunas do modelo "Adicional Noturno Professores" pelo
    cabeçalho (linha 1): Matrícula, Professor, Adicional Noturno, Horas/Aula,
    Adicional Noturno (Minutos) e Unidade. Retorna {chave: índice} ou None
    quando falta algum cabeçalho obrigatório. Nomes que só contêm a palavra
    ('Minutos Decimais', 'Adicional Noturno (Horas)') não confundem: minutos
    é a primeira coluna terminada em '(minutos)' e horas é 'Horas/Aula'."""
    cabecalho = next(aba.iter_rows(min_row=1, max_row=1, values_only=True), ())
    colunas = {}

    def _mapear(chave, indice):
        colunas.setdefault(chave, indice)

    for indice, valor in enumerate(cabecalho):
        nome = _normalizar(valor)
        if not nome:
            continue
        if nome == 'matricula':
            _mapear('matricula', indice)
        elif nome in ('professor', 'nome'):
            _mapear('professor', indice)
        elif nome == 'adicional noturno':
            _mapear('adicional', indice)
        elif nome == 'horas/aula':
            _mapear('horas', indice)
        elif nome.endswith('(minutos)'):
            _mapear('minutos', indice)
        elif nome.startswith('unidade'):
            _mapear('unidade', indice)
    obrigatorias = ('matricula', 'professor', 'adicional', 'minutos')
    return colunas if all(chave in colunas for chave in obrigatorias) else None


def _linha_do_modelo(valores, colunas):
    """Linha da planilha → dicionário no formato dos lançamentos; None quando
    a linha está totalmente vazia. Campos ausentes/ilegíveis ficam None
    (horas/unidade) ou 0 (minutos), como na digitação manual."""
    def _texto(chave):
        indice = colunas.get(chave)
        if indice is None or indice >= len(valores):
            return None
        valor = valores[indice]
        if valor is None:
            return None
        return str(valor).strip() or None

    if all(v is None for v in valores):
        return None
    matricula = _texto('matricula')
    # Matrícula numérica lida como float vem '13164.0'.
    if matricula and re.fullmatch(r'\d+\.0', matricula):
        matricula = matricula[:-2]
    horas_brutas = _texto('horas')
    try:
        horas = float(horas_brutas.replace(',', '.')) if horas_brutas else None
    except ValueError:
        horas = None
    minutos_brutos = _texto('minutos')
    try:
        minutos = int(float(minutos_brutos)) if minutos_brutos else 0
    except ValueError:
        minutos = 0
    unidade_bruta = _texto('unidade')
    try:
        unidade = int(float(unidade_bruta)) if unidade_bruta else None
    except ValueError:
        unidade = None
    return {
        'registration': matricula,
        'full_name': _texto('professor'),
        'entitled': (_texto('adicional') or '').lower() == 'sim',
        'class_hours': horas,
        'night_minutes': minutos,
        'uo_code': unidade,
    }


# ── Rotas ────────────────────────────────────────────────────────────────────

@bp.route('/')
@login_required
@require_permission('an:read')
@require_module('finance')
def registros():
    """Adicional Noturno - Professores: formulário de lançamento em cima e a
    listagem (com filtro por mês) abaixo, na mesma página."""
    form = _preparar_form(FormAdicionalNoturno())
    return _render_registros(form,
                             month_filter=_mes_valido(request.args.get('month_base')))


@bp.route('/adicionar', methods=['POST'])
@login_required
@require_permission('an:create')
@require_module('finance')
def adicionar():
    form = _preparar_form(FormAdicionalNoturno())
    if form.validate_on_submit():
        try:
            full_name, registration = _resolver_professor(form)
        except ValueError as exc:
            flash(str(exc), 'warning')
            return _render_registros(
                form,
                month_filter=_mes_valido(request.args.get('month_base')))
        db.session.add(TeacherNightShiftPay(
            month_base=form.month_base.data,
            registration=registration,
            full_name=full_name,
            entitled=form.entitled.data == 'Sim',
            class_hours=_class_hours_decimal(form.class_hours.data),
            night_minutes=form.night_minutes.data,
            uo_code=form.uo_code.data,
            unity_id=current_unity_id(), created_by_id=current_user.id))
        db.session.commit()
        flash(f'Lançamento adicionado: {full_name} — '
              f'{form.night_minutes.data} minuto(s) noturno(s) em '
              f'{_rotulo_mes(form.month_base.data)}.', 'success')
        # Mantém o filtro ativo: quem lançou filtrando por um mês continua
        # vendo a lista do mês.
        return redirect_preserving_args('an.registros')
    return _render_registros(
        form,
        month_filter=_mes_valido(request.args.get('month_base')))


@bp.route('/importar', methods=['POST'])
@login_required
@require_permission('an:create')
@require_module('finance')
def importar():
    """Importa a planilha "Adicional Noturno Professores" (.xlsx): localiza
    as colunas pelo cabeçalho da primeira linha e cria um lançamento por
    linha com matrícula e professor — additional "Não" e minutos 0 entram
    como estão (a exportação aplica os critérios do gerador). Quem já possui
    lançamento no mês selecionado é mantido, então reimportar não duplica."""
    arquivo = request.files.get('planilha')
    if not arquivo or not arquivo.filename:
        flash('Selecione a planilha (.xlsx) para importar.', 'warning')
        return redirect(url_for('an.registros'))
    if not arquivo.filename.lower().endswith('.xlsx'):
        flash('Formato não suportado: envie a planilha em .xlsx.', 'danger')
        return redirect(url_for('an.registros'))

    mes = _mes_valido(request.form.get('month_base')) or _mes_atual()

    try:
        planilha = load_workbook(arquivo, data_only=True, read_only=True)
    except Exception:
        flash('Não foi possível ler o arquivo — confira se é uma planilha '
              '.xlsx válida.', 'danger')
        return redirect(url_for('an.registros'))

    aba = planilha.worksheets[0]
    colunas = _colunas_do_modelo(aba)
    if colunas is None:
        flash('A planilha não segue o modelo "Adicional Noturno Professores" '
              '(linha 1 com os cabeçalhos Matrícula, Professor, Adicional '
              'Noturno, Horas/Aula, Adicional Noturno (Minutos) e Unidade).',
              'danger')
        return redirect(url_for('an.registros'))

    existentes = {r for (r,) in db.session.query(TeacherNightShiftPay.registration)
                  .filter_by(unity_id=current_unity_id(), month_base=mes).all()}
    criados, mantidos, ignoradas = 0, [], 0
    for valores in aba.iter_rows(min_row=2, values_only=True):
        dados = _linha_do_modelo(valores, colunas)
        if dados is None:
            continue  # linha em branco da planilha
        if not dados['registration'] or not dados['full_name']:
            ignoradas += 1
            continue
        if dados['registration'] in existentes:
            mantidos.append(dados['full_name'])
            continue
        db.session.add(TeacherNightShiftPay(
            month_base=mes,
            registration=dados['registration'][:20],
            full_name=dados['full_name'][:255],
            entitled=dados['entitled'],
            class_hours=dados['class_hours'],
            night_minutes=dados['night_minutes'],
            uo_code=dados['uo_code'],
            unity_id=current_unity_id(), created_by_id=current_user.id))
        existentes.add(dados['registration'])
        criados += 1
    db.session.commit()

    if criados:
        flash(f'{criados} lançamento(s) importado(s) para '
              f'{_rotulo_mes(mes)}.', 'success')
    else:
        flash(f'Nenhum lançamento novo para {_rotulo_mes(mes)} — a planilha '
              'não trouxe linhas com matrícula e professor.', 'warning')
    if mantidos:
        flash(f'{len(mantidos)} professor(es) já possuíam lançamento no mês e '
              f'foram mantidos: {", ".join(mantidos)}.', 'info')
    if ignoradas:
        flash(f'{ignoradas} linha(s) sem matrícula ou professor foram '
              'ignoradas.', 'info')
    # Abre no mês da importação: o operador confere o que acabou de entrar.
    return redirect(url_for('an.registros', month_base=mes))


@bp.route('/<int:registro_id>/editar', methods=['GET', 'POST'])
@login_required
@require_permission('an:edit')
@require_module('finance')
def editar(registro_id):
    registro = _get_registro_scoped(registro_id)
    form = _preparar_form(FormAdicionalNoturno(obj=registro))

    if request.method == 'GET':
        # O campo entitled é Sim/Não no formulário e booleano no banco.
        form.entitled.data = 'Sim' if registro.entitled else 'Não'
        # Lançamento sem horas/aula gravadas: campo em branco.
        form.class_hours.data = (f'{registro.class_hours:.2f}'.replace('.', ',')
                                 if registro.class_hours is not None else '')
        form.month_base.data = registro.month_base
        form.registration.data = registro.registration
        # Pré-seleciona o professor do cadastro com a mesma matrícula; sem
        # correspondência, a edição abre no modo manual com o nome gravado.
        professor_id = _professor_id_por_matricula(registro.registration)
        if professor_id:
            form.teacher.data = professor_id
        else:
            form.teacher.data = MANUAL
            form.full_name.data = registro.full_name

    if form.validate_on_submit():
        try:
            full_name, registration = _resolver_professor(form)
        except ValueError as exc:
            flash(str(exc), 'warning')
            return render_template('an/registro_form.html', form=form,
                                   registro=registro)
        registro.month_base = form.month_base.data
        registro.registration = registration
        registro.full_name = full_name
        registro.entitled = form.entitled.data == 'Sim'
        registro.class_hours = _class_hours_decimal(form.class_hours.data)
        registro.night_minutes = form.night_minutes.data
        registro.uo_code = form.uo_code.data
        db.session.commit()
        flash('Alteração realizada!', 'success')
        return redirect_preserving_args('an.registros')

    return render_template('an/registro_form.html', form=form, registro=registro,
                           professores_map=_professores_map())


@bp.route('/<int:registro_id>/excluir', methods=['POST'])
@login_required
@require_permission('an:delete')
@require_module('finance')
def excluir(registro_id):
    registro = _get_registro_scoped(registro_id)
    db.session.delete(registro)
    db.session.commit()
    flash('Lançamento do Adicional Noturno excluído.', 'success')
    return redirect_back('an.registros')


@bp.route('/excluir-todos', methods=['POST'])
@login_required
@require_permission('an:delete')
@require_module('finance')
def excluir_todos():
    """Remove TODOS os lançamentos de Adicional Noturno da unidade atual,
    independentemente do filtro aplicado na listagem."""
    qtde = TeacherNightShiftPay.query \
        .filter_by(unity_id=current_unity_id()).delete()
    db.session.commit()
    flash(f'{qtde} lançamento(s) do Adicional Noturno removido(s).', 'success')
    return redirect(url_for('an.registros'))


@bp.route('/exportar')
@login_required
@require_permission('an:export')
@require_module('finance')
def exportar():
    """Gera a planilha de pagamento a partir dos lançamentos da unidade
    ativa (do mês filtrado, quando houver), no modelo planilha_base_an.xlsx
    — Matrícula, Nome, H/A, Minutos Decimais e UO Lotação a partir da linha
    5, nos mesmos moldes do gerador original (somente adicional "Sim" com
    minutos maiores que zero, ordenado pelo código da unidade)."""
    month_filter = _mes_valido(request.args.get('month_base'))
    query = TeacherNightShiftPay.query.filter_by(unity_id=current_unity_id())
    if month_filter:
        query = query.filter_by(month_base=month_filter)
    registros = [r for r in query.all() if r.is_exportable]
    # Ordenação do gerador original: pelo código da unidade, mantendo dentro
    # de cada unidade a ordem de lançamento (a ordem das linhas da planilha
    # de origem, no script).
    registros.sort(key=lambda r: (r.uo_code or 0, r.id))

    if not registros:
        flash('Nenhum lançamento elegível para exportar — é preciso ter '
              'Adicional Noturno "Sim" e minutos maiores que 0.', 'warning')
        return redirect(url_for('an.registros', month_base=month_filter))

    template_path = os.path.join(current_app.root_path, 'static',
                                 'templates_excel', AN_TEMPLATE_FILE)
    try:
        workbook = load_workbook(template_path)
    except FileNotFoundError:
        flash(f'Modelo do Excel não encontrado no servidor ({AN_TEMPLATE_FILE}).',
              'danger')
        return redirect(url_for('an.registros'))

    worksheet = workbook[AN_SOURCE_SHEET]

    # Somente as colunas de dados (A, B, C, G e K) recebem valores — todas as
    # demais seguem com as fórmulas do modelo fornecido (dias, totais e valor
    # do adicional), inclusive a aba oculta "Não mexer nesta planilha", que
    # referencia estas linhas. fullCalcOnLoad garante o recálculo das fórmulas
    # na abertura do arquivo gerado.
    workbook.calculation.fullCalcOnLoad = True

    truncated = len(registros) > EXPORT_MAX_ROWS
    larguras = {}

    def _medir(letra, texto):
        larguras[letra] = max(larguras.get(letra, 0), len(texto))

    for offset, registro in enumerate(registros[:EXPORT_MAX_ROWS]):
        row = EXPORT_FIRST_ROW + offset
        matricula = registro.registration
        # O modelo espera a matrícula numérica.
        if isinstance(matricula, str) and matricula.isdigit():
            matricula = int(matricula)
        worksheet.cell(row=row, column=1, value=matricula)                  # A
        worksheet.cell(row=row, column=2, value=registro.full_name)         # B
        horas = (float(registro.class_hours)
                 if registro.class_hours is not None else 0.0)
        if registro.class_hours is not None:                                # C
            worksheet.cell(row=row, column=3, value=horas)
        decimais = registro.decimal_minutes
        worksheet.cell(row=row, column=7, value=decimais)                   # G
        worksheet.cell(row=row, column=11, value=registro.uo_code)          # K

        # Mede também as colunas de fórmula (mesmo cálculo do cache) para
        # alargar o que o modelo estreito não exibe (E viraria '###').
        d = float(worksheet.cell(row=row, column=4).value or 0)
        dias = float(worksheet.cell(row=row, column=6).value or 0)
        g2 = float(worksheet['G2'].value or 0)
        e = horas * d
        h = decimais / g2 if g2 else 0.0
        i = h * dias
        _medir('A', str(matricula))
        _medir('B', registro.full_name)
        _medir('C', f'{horas:.2f}')
        _medir('E', f'{e:.2f}')
        _medir('G', f'{decimais:.2f}')
        _medir('H', f'{h:.4f}')
        _medir('I', f'{i:.2f}')
        _medir('J', f'{e * i:.2f}')
        _medir('K', str(registro.uo_code))

    _ajustar_larguras(worksheet, larguras)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    # Repõe os resultados em cache das fórmulas do modelo (descartados pelo
    # openpyxl) para as colunas calculadas aparecerem ao abrir o arquivo,
    # mesmo no Modo de Exibição Protegido — as fórmulas seguem intactas.
    conteudo = _preencher_cache_formulas(output.getvalue())

    if truncated:
        flash(f'Atenção: o modelo suporta {EXPORT_MAX_ROWS} linhas — a '
              f'exportação incluiu apenas as {EXPORT_MAX_ROWS} primeiras em '
              'ordem de unidade.', 'warning')

    stamp = datetime.now().strftime('%d-%m-%Y %H-%M-%S')
    filename = f'Tabela Adicional Noturno Professores ({stamp}).xlsx'
    return send_file(io.BytesIO(conteudo), as_attachment=True,
                     download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument'
                              '.spreadsheetml.sheet')
